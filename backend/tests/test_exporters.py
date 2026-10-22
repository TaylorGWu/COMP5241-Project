from __future__ import annotations

import io

from pptx import Presentation

from app.exporters.pptx import slide_deck_to_pptx
from app.schemas.ir import Project, ProjectMeta, Slide, SlideDeck


def test_pptx_export_ignores_optional_animation_metadata() -> None:
    """The exporter is a pure static IR transform, not a media renderer."""
    project = Project(id="p1", meta=ProjectMeta(title="Git", audience="Beginners", duration_min=30))
    deck = SlideDeck(
        artifact_id="a-slides",
        slides=[
            Slide(
                id="s1",
                order=1,
                layout="bullets",
                title="Record a checkpoint",
                bullets=["A commit creates a recoverable history entry."],
                animation={
                    "steps": [
                        {"kind": "reveal", "target": "title", "atMs": 0},
                        {"kind": "emphasis", "target": "bullet:0", "atMs": 700},
                    ]
                },
            )
        ],
    )

    exported = slide_deck_to_pptx(project, deck)
    output = Presentation(io.BytesIO(exported))
    text = "\n".join(
        shape.text
        for shape in output.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert len(output.slides) == 1
    assert "Record a checkpoint" in text
    assert "recoverable history entry" in text
    assert "reveal" not in text and "emphasis" not in text
