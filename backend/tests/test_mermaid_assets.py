"""Generated diagrams must be real PNG assets in the matching chapter deck."""

from __future__ import annotations

import io

import httpx
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.api.routes import demo
from app.main import app
from app.services import mermaid_assets
from app.services.mermaid_assets import mermaid_render_config, render_mermaid_png
from app.services.svg_diagrams import render_course_svg


@pytest.fixture(autouse=True)
def reset_courses(monkeypatch: pytest.MonkeyPatch):
    demo.reset_course_workspace()
    monkeypatch.setattr(demo, "FAKE_MODEL_DELAY", 0.0)


async def request(method: str, path: str, body: dict | None = None):
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        return await client.request(method, path, json=body)


def test_diagrams_use_the_course_visual_theme_and_generous_flow_spacing():
    config = mermaid_render_config()

    assert config["theme"] == "base"
    assert config["themeVariables"]["primaryColor"] == "#eef4ff"
    assert config["themeVariables"]["lineColor"] == "#5472a6"
    assert config["flowchart"]["nodeSpacing"] >= 48
    assert config["flowchart"]["rankSpacing"] >= 56


def test_diagram_png_has_a_deterministic_fallback_when_browser_renderer_is_unavailable(
    monkeypatch: pytest.MonkeyPatch,
):
    """PPTX export must still contain a visual without a Puppeteer browser."""
    monkeypatch.setattr(mermaid_assets, "_cli", lambda: "mmdc")
    monkeypatch.setattr(
        mermaid_assets.subprocess,
        "run",
        lambda *_args, **_kwargs: type("Result", (), {"returncode": 1})(),
    )
    render_mermaid_png.cache_clear()

    image = render_mermaid_png("graph LR; A[Change]-->B[Review]-->C[Verified result]")

    assert image.startswith(b"\x89PNG\r\n\x1a\n")
    with Image.open(io.BytesIO(image)) as rendered:
        assert rendered.size == (1920, 1080)


def test_png_fallback_uses_the_same_semantic_template_as_the_svg_preview():
    decision = "graph LR\nA[git pull] --> B{同一区域?}\nB -- 否 --> C[自动合并]\nB -- 是 --> D[手动解决]"
    exchange = "graph LR\nA[开发者A 本地] -->|push| R[远程仓库]\nR -->|pull| B[开发者B 本地]"

    assert mermaid_assets.fallback_template_for(decision) == "decision-branch"
    assert mermaid_assets.fallback_template_for(exchange) == "exchange-lane"


def test_course_svg_is_semantic_and_escapes_mermaid_labels():
    svg = render_course_svg(
        "flowchart LR\nmain[主幹] --> feature[建立分支] --> review[程式碼審查] --> merge[合併]"
    )

    assert svg.startswith("<svg")
    assert 'data-template="branch-lane"' in svg
    assert "建立分支" in svg
    assert "<script" not in svg


def test_course_svg_uses_a_decision_template_when_mermaid_has_yes_no_branches():
    svg = render_course_svg(
        "graph LR\nA[git pull] --> B{同一区域?}\nB -- 否 --> C[自动合并]\nB -- 是 --> D[手动解决]"
    )

    assert 'data-template="decision-branch"' in svg
    assert "自动合并" in svg
    assert "手动解决" in svg


def test_course_svg_uses_exchange_lanes_for_local_remote_sync():
    svg = render_course_svg(
        "graph LR\nA[开发者A 本地] -->|push| R[远程仓库]\nR -->|pull| B[开发者B 本地]"
    )

    assert 'data-template="exchange-lane"' in svg
    assert "push" in svg
    assert "pull" in svg


def test_course_svg_never_invents_numbered_placeholder_steps_for_a_short_graph():
    svg = render_course_svg("graph LR; A[核心机制] --> B[应用]")

    assert "核心机制" in svg
    assert "应用" in svg
    assert "Step 1" not in svg
    assert "Step 2" not in svg


async def test_generated_chapter_exports_an_embedded_diagram_image():
    created = await request("POST", "/api/demo/courses", {"title": "Git", "audience": "Beginners", "durationMin": 30})
    course = created.json()
    chapter_id = course["chapterIds"][0]
    generated = await request(
        "POST",
        f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/generate",
        {"topic": "Git branches", "audience": "Beginners", "duration_min": 30},
    )
    assert generated.status_code == 200, generated.text
    assert all(asset["provider"] == "mermaid-cli" for asset in generated.json()["package"]["assets"] if asset["kind"] == "diagram")

    exported = await request("GET", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/export/deck.pptx")

    assert exported.status_code == 200, exported.text
    deck = Presentation(io.BytesIO(exported.content))
    assert len(deck.slides) == len(generated.json()["package"]["slideDeck"]["slides"])
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pictures
    text_runs = [run for slide in deck.slides for shape in slide.shapes if shape.has_text_frame for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text]
    assert text_runs
    assert all(run.font.name == "Hiragino Sans GB" for run in text_runs)
    assert all(run._r.rPr.xpath("a:ea") for run in text_runs)
    assert all(picture.left >= 0 and picture.top >= 0 for picture in pictures)
    assert all(picture.left + picture.width <= deck.slide_width for picture in pictures)
    assert all(picture.top + picture.height <= deck.slide_height for picture in pictures)


async def test_chapter_diagram_endpoint_returns_png_from_its_asset_id():
    created = await request("POST", "/api/demo/courses", {"title": "Git", "audience": "Beginners", "durationMin": 30})
    course = created.json()
    chapter_id = course["chapterIds"][0]
    generated = await request("POST", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/generate", {"topic": "Git branches", "audience": "Beginners", "duration_min": 30})
    assert generated.status_code == 200, generated.text
    asset_id = generated.json()["package"]["assets"][0]["id"]

    image = await request("GET", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/assets/{asset_id}.png")

    assert image.status_code == 200, image.text
    assert image.headers["content-type"].startswith("image/png")
    assert image.headers["cache-control"] == "no-store"
    assert image.content.startswith(b"\x89PNG\r\n\x1a\n")


async def test_chapter_diagram_endpoint_returns_semantic_svg_from_its_asset_id():
    created = await request("POST", "/api/demo/courses", {"title": "Git", "audience": "Beginners", "durationMin": 30})
    course = created.json()
    chapter_id = course["chapterIds"][0]
    generated = await request("POST", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/generate", {"topic": "Git branches", "audience": "Beginners", "duration_min": 30})
    assert generated.status_code == 200, generated.text
    asset_id = generated.json()["package"]["assets"][0]["id"]

    image = await request("GET", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/assets/{asset_id}.svg")

    assert image.status_code == 200, image.text
    assert image.headers["content-type"].startswith("image/svg+xml")
    assert image.headers["cache-control"] == "no-store"
    assert 'data-template=' in image.text


async def test_deck_export_can_append_question_set_without_mutating_slides():
    created = await request("POST", "/api/demo/courses", {"title": "Git", "audience": "Beginners", "durationMin": 30})
    course = created.json()
    chapter_id = course["chapterIds"][0]
    generated = await request("POST", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/generate", {"topic": "Git branches", "audience": "Beginners", "duration_min": 30})
    assert generated.status_code == 200, generated.text

    questions = await request("POST", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/questions", {"count_per_objective": 1, "language": "zh-HK"})
    assert questions.status_code == 200, questions.text
    package = questions.json()["package"]
    source_count = len(package["slideDeck"]["slides"])
    question_count = len(package["questionSet"]["questions"])

    exported = await request("GET", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/export/deck.pptx?include_questions=true")

    assert exported.status_code == 200, exported.text
    deck = Presentation(io.BytesIO(exported.content))
    assert len(deck.slides) == source_count + question_count
    assert len(questions.json()["package"]["slideDeck"]["slides"]) == source_count
    deck_slides = list(deck.slides)
    notes = [slide.notes_slide.notes_text_frame.text for slide in deck_slides[-question_count:]]
    assert all("答案：" in note and "解析：" in note for note in notes)


async def test_tall_diagram_is_contained_within_slide(monkeypatch: pytest.MonkeyPatch):
    tall = io.BytesIO()
    Image.new("RGB", (100, 1200), "white").save(tall, "PNG")
    monkeypatch.setattr(demo, "render_mermaid_png", lambda _code: tall.getvalue())

    created = await request("POST", "/api/demo/courses", {"title": "Git", "audience": "Beginners", "durationMin": 30})
    course = created.json()
    chapter_id = course["chapterIds"][0]
    generated = await request("POST", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/generate", {"topic": "Git branches", "audience": "Beginners", "duration_min": 30})
    assert generated.status_code == 200, generated.text

    exported = await request("GET", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/export/deck.pptx")
    assert exported.status_code == 200, exported.text
    deck = Presentation(io.BytesIO(exported.content))
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pictures
    assert all(picture.top + picture.height <= deck.slide_height for picture in pictures)


async def test_one_failed_diagram_does_not_abort_deck(monkeypatch: pytest.MonkeyPatch):
    good = io.BytesIO()
    Image.new("RGB", (300, 100), "white").save(good, "PNG")

    def render(code: str) -> bytes:
        if "核心机制" in code:
            raise RuntimeError("unexpected renderer failure")
        return good.getvalue()

    monkeypatch.setattr(demo, "render_mermaid_png", render)
    created = await request("POST", "/api/demo/courses", {"title": "Git", "audience": "Beginners", "durationMin": 30})
    course = created.json()
    chapter_id = course["chapterIds"][0]
    generated = await request("POST", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/generate", {"topic": "Git branches", "audience": "Beginners", "duration_min": 30})
    assert generated.status_code == 200, generated.text

    exported = await request("GET", f"/api/demo/courses/{course['id']}/chapters/{chapter_id}/export/deck.pptx")
    assert exported.status_code == 200, exported.text
    deck = Presentation(io.BytesIO(exported.content))
    picture_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    assert 0 < picture_count < len(deck.slides)
    assert any("尚未渲染" in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
