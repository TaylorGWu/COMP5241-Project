"""课件 → pptx —— 设计规格 §12.2。

═══════════════════════════════════════════════════════════════════
版式:一页 IR 幻灯片 = 一页 pptx,不做别的
═══════════════════════════════════════════════════════════════════
不额外加封面页、不加结束页。原因很实在:教师改课件之后重新导出,
页数必须还对得上,否则「课件第 3 页」这句话在两个地方指的是不同的页。

封面信息(受众、课时)挂在**第一页**上,而不是单独占一页。

═══════════════════════════════════════════════════════════════════
版面:固定 6 种,不追求任意排版
═══════════════════════════════════════════════════════════════════
title / bullets / diagram / twoCol / code / quiz —— 与 IR 的 `SlideLayout`
一一对应。规格里写明了不追求任意排版:占位符 + 自适应文本框极易溢出,
而溢出的表现是文字跑到页面外,教师拿到手才发现。

═══════════════════════════════════════════════════════════════════
图:插不进去也要留占位
═══════════════════════════════════════════════════════════════════
规格要求图用后端渲染的 PNG。`slide_deck_to_pptx(..., asset_paths=...)` 时:
  有图 → 真的插进去
  没图 → 放一个占位框,写明这是什么图

**绝不留空白。** 一个空框会让教师以为导出坏了,而实际上是渲染还没接上 ——
这两种情况的处理方式完全不同。
"""

from __future__ import annotations

import io
from collections.abc import Mapping
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from app.schemas.ir import Project, Question, QuestionSet, Slide, SlideDeck, Visual

#: 16:9。和网页预览的画布比例一致 —— 不然教师会以为放错了。
SLIDE_WIDTH = 12192000
SLIDE_HEIGHT = 6858000

#: 页面边距(EMU)
MARGIN = 685800

#: 各版面里正文区的位置(precise 一点比用 layout placeholder 稳)
TITLE_TOP = 2286000
BODY_TOP = 1828800
BODY_HEIGHT = 4114800

FONT_TITLE = Pt(40)
FONT_SLIDE_TITLE = Pt(28)
FONT_BODY = Pt(18)
FONT_META = Pt(12)
FONT_CODE = Pt(14)
FONT_FAMILY = "Hiragino Sans GB"

INK = "172033"
CANVAS = "F5F7FB"
SURFACE = "FFFFFF"
BLUE = "315FC6"
BLUE_SOFT = "EAF1FF"
BODY = "465774"

#: 图的占位框。**一定要带「图」字** —— 界面上和测试里都靠它区分「有图」和「没接渲染」。
PLACEHOLDER_PREFIX = "【图】"


def _blank_slide(prs: Presentation):
    """空白版面 —— 所有内容都由我们自己摆位置。"""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, *, left: int, top: int, width: int, height: int):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def _fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def _panel(slide, *, left: int, top: int, width: int, height: int, color: str = SURFACE, rounded: bool = True):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Emu(left), Emu(top), Emu(width), Emu(height),
    )
    _fill(shape, color)
    return shape


def _background(slide, color: str) -> None:
    _panel(slide, left=0, top=0, width=SLIDE_WIDTH, height=SLIDE_HEIGHT, color=color, rounded=False)


def _line(frame, text: str, *, size, bold: bool = False, color: str = INK) -> None:
    """往文本框追加一行。

    python-pptx 的 text_frame 永远自带一个空段落,所以第一行要**复用**它,
    否则页面上会多出一个空行(而且是行首的空行,最显眼的位置)。
    """
    paragraph = frame.paragraphs[0] if not frame.paragraphs[0].text else frame.add_paragraph()
    paragraph.text = text
    paragraph.font.size = size
    paragraph.font.bold = bold
    paragraph.font.name = FONT_FAMILY
    paragraph.font.color.rgb = RGBColor.from_string(color)
    for run in paragraph.runs:
        run.font.name = FONT_FAMILY
        run.font.color.rgb = RGBColor.from_string(color)
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", FONT_FAMILY)
        run._r.get_or_add_rPr().append(east_asian)


def _visual_placeholder(visual: Visual) -> str:
    """图插不进去时显示的文字。

    Mermaid 的源码对教师也有用 —— 至少能看出这里该有张什么图。
    """
    kind = {"diagram": "示意图", "chart": "图表", "icon": "图标", "illustration": "插画"}.get(
        visual.kind, visual.kind
    )
    code = (visual.spec or {}).get("code") or (visual.spec or {}).get("prompt") or ""
    hint = str(code).strip().replace("\n", " ")[:120]
    return f"{PLACEHOLDER_PREFIX}{kind}(尚未渲染){(':' + hint) if hint else ''}"


def _add_visual(
    slide,
    visual: Visual,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    asset_paths: Mapping[str, Path] | None,
) -> None:
    """把一处视觉需求落到页面上:能插图就插,不能就留占位。"""
    path = asset_paths.get(visual.asset_id) if (asset_paths and visual.asset_id) else None
    if path is not None and Path(path).exists():
        with Image.open(path) as image:
            image_width, image_height = image.size
        scale = min(width / image_width, height / image_height)
        picture_width = round(image_width * scale)
        picture_height = round(image_height * scale)
        slide.shapes.add_picture(
            str(path),
            Emu(left + (width - picture_width) // 2),
            Emu(top + (height - picture_height) // 2),
            width=Emu(picture_width),
            height=Emu(picture_height),
        )
        return

    frame = _textbox(slide, left=left, top=top, width=width, height=min(height, 914400))
    frame.paragraphs[0].text = _visual_placeholder(visual)
    frame.paragraphs[0].font.size = FONT_META
    frame.paragraphs[0].font.name = FONT_FAMILY


def _render_slide(
    prs: Presentation,
    slide_ir: Slide,
    *,
    is_first: bool,
    meta_line: str,
    asset_paths: Mapping[str, Path] | None,
) -> None:
    """把一页 IR 渲染成一页 pptx。"""
    slide = _blank_slide(prs)

    if slide_ir.layout == "title":
        _background(slide, INK)
        _panel(slide, left=MARGIN, top=MARGIN, width=1143000, height=411480, color=BLUE, rounded=True)
        label = _textbox(slide, left=MARGIN + 182880, top=MARGIN + 68580, width=777240, height=274320)
        _line(label, "COURSE NOTE", size=FONT_META, bold=True, color=SURFACE)
        frame = _textbox(
            slide,
            left=MARGIN,
            top=1905000,
            width=SLIDE_WIDTH - 2 * MARGIN,
            height=1600200,
        )
        _line(frame, slide_ir.title, size=FONT_TITLE, bold=True, color=SURFACE)
        for bullet in slide_ir.bullets:
            _line(frame, bullet, size=FONT_BODY, color="DCE8FF")
        if is_first and meta_line:
            # 封面信息挂在这一页上,不单独占一页 —— 否则教师改完课件重新导出,
            # 页数就对不上了,而「课件第 3 页」这句话在两个地方会指不同的页。
            _line(frame, meta_line, size=FONT_META, color="AFC5F7")
        slide.notes_slide.notes_text_frame.text = slide_ir.speaker_notes
        return

    _background(slide, CANVAS)
    _panel(slide, left=MARGIN, top=MARGIN, width=914400, height=342900, color=BLUE_SOFT, rounded=True)
    tag = _textbox(slide, left=MARGIN + 137160, top=MARGIN + 45720, width=640080, height=228600)
    _line(tag, slide_ir.layout.upper(), size=FONT_META, bold=True, color=BLUE)

    # 除封面外,第一行都是标题
    title_frame = _textbox(
        slide, left=MARGIN, top=1143000, width=SLIDE_WIDTH - 2 * MARGIN, height=868680
    )
    _line(title_frame, slide_ir.title, size=FONT_SLIDE_TITLE, bold=True)

    if is_first and meta_line:
        meta_frame = _textbox(
            slide, left=MARGIN, top=1935480, width=SLIDE_WIDTH - 2 * MARGIN, height=274320
        )
        _line(meta_frame, meta_line, size=FONT_META)

    has_visual = bool(slide_ir.visuals)
    body_width = SLIDE_WIDTH - 2 * MARGIN
    if has_visual:
        # 图是主视觉;正文只承担解释,不再和图五五分抢空间。
        body_width = int((SLIDE_WIDTH - 2 * MARGIN) * 0.36)

    if has_visual:
        _panel(slide, left=MARGIN, top=BODY_TOP, width=body_width, height=BODY_HEIGHT)
        visual_left = MARGIN + body_width + 228600
        _panel(slide, left=visual_left, top=BODY_TOP, width=SLIDE_WIDTH - MARGIN - visual_left, height=BODY_HEIGHT, color=SURFACE)
    elif slide_ir.layout != "twoCol":
        _panel(slide, left=MARGIN, top=BODY_TOP, width=body_width, height=BODY_HEIGHT)

    if slide_ir.layout == "twoCol":
        # 左右分栏:奇数条左边,偶数条右边
        left_items = slide_ir.bullets[::2]
        right_items = slide_ir.bullets[1::2]
        half = int((SLIDE_WIDTH - 3 * MARGIN) / 2)
        _panel(slide, left=MARGIN, top=BODY_TOP, width=half, height=BODY_HEIGHT)
        _panel(slide, left=MARGIN * 2 + half, top=BODY_TOP, width=half, height=BODY_HEIGHT)
        left_frame = _textbox(slide, left=MARGIN, top=BODY_TOP, width=half, height=BODY_HEIGHT)
        for item in left_items:
            _line(left_frame, item, size=FONT_BODY)
        right_frame = _textbox(
            slide, left=MARGIN * 2 + half, top=BODY_TOP, width=half, height=BODY_HEIGHT
        )
        for item in right_items:
            _line(right_frame, item, size=FONT_BODY)
    elif slide_ir.bullets:
        body_frame = _textbox(
            slide, left=MARGIN, top=BODY_TOP, width=body_width, height=BODY_HEIGHT
        )
        size = FONT_CODE if slide_ir.layout == "code" else FONT_BODY
        for bullet in slide_ir.bullets:
            _line(body_frame, bullet, size=size)

    if has_visual:
        visual = slide_ir.visuals[0]
        _add_visual(
            slide,
            visual,
            left=MARGIN + body_width + 228600,
            top=BODY_TOP + 137160,
            width=SLIDE_WIDTH - 2 * MARGIN - body_width - 228600,
            height=BODY_HEIGHT - 274320,
            asset_paths=asset_paths,
        )

    slide.notes_slide.notes_text_frame.text = slide_ir.speaker_notes


def _render_question_slide(prs: Presentation, question: Question, *, number: int) -> None:
    """Append a teacher-friendly practice page without mutating the source deck."""
    slide = _blank_slide(prs)
    _background(slide, CANVAS)
    _panel(slide, left=MARGIN, top=MARGIN, width=1543050, height=342900, color=BLUE_SOFT, rounded=True)
    tag = _textbox(slide, left=MARGIN + 137160, top=MARGIN + 45720, width=1257300, height=228600)
    _line(tag, f"PRACTICE {number:02d}", size=FONT_META, bold=True, color=BLUE)
    title = _textbox(slide, left=MARGIN, top=1028700, width=SLIDE_WIDTH - 2 * MARGIN, height=1219200)
    _line(title, question.stem, size=FONT_SLIDE_TITLE, bold=True)
    option_top = 2514600
    for index, option in enumerate(question.options):
        column = index % 2
        row = index // 2
        width = int((SLIDE_WIDTH - 3 * MARGIN) / 2)
        left = MARGIN + column * (width + MARGIN)
        top = option_top + row * 914400
        _panel(slide, left=left, top=top, width=width, height=640080, color=SURFACE)
        frame = _textbox(slide, left=left + 182880, top=top + 137160, width=width - 365760, height=365760)
        _line(frame, f"{option.id}. {option.text}", size=FONT_BODY)
    answer = next((option.text for option in question.options if option.id == question.answer), question.answer)
    slide.notes_slide.notes_text_frame.text = f"答案：{answer}\n解析：{question.explanation}"


def slide_deck_to_pptx(
    project: Project,
    deck: SlideDeck,
    *,
    asset_paths: Mapping[str, Path] | None = None,
    question_set: QuestionSet | None = None,
    include_questions: bool = False,
) -> bytes:
    """把课件导出成 pptx,返回文件字节。

    `asset_paths` 是 `assetId → 本地 PNG 路径`。给不出来就走占位 ——
    见模块文档字符串里关于「绝不留空白」的说明。
    """
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH)
    prs.slide_height = Emu(SLIDE_HEIGHT)

    meta_line = f"受众:{project.meta.audience}　课时:{project.meta.duration_min} 分钟"

    for index, slide_ir in enumerate(deck.slides):
        _render_slide(
            prs,
            slide_ir,
            is_first=(index == 0),
            meta_line=meta_line,
            asset_paths=asset_paths,
        )

    if include_questions and question_set is not None:
        for index, question in enumerate(question_set.questions, start=1):
            _render_question_slide(prs, question, number=index)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
