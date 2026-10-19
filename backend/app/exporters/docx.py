"""教案与考卷 → docx —— 设计规格 §12.2。

═══════════════════════════════════════════════════════════════════
为什么教案和考卷都用 docx
═══════════════════════════════════════════════════════════════════
规格里写的是「docx | 教案、考卷」。两者都是**给人读、给人改**的东西:
教师要在教案上批注,要在考卷上改题。pptx 不适合承载成段的文字和表格。

═══════════════════════════════════════════════════════════════════
答案卷做成开关,不做成两份模板
═══════════════════════════════════════════════════════════════════
教师版和学生版的区别只有「带不带答案和解析」。做成两份模板的话,
迟早会漂移 —— 改了一处忘了另一处,学生就拿不到最新的题。

所以是同一个函数加一个 `include_answer_key` 参数。

**解析只在答案区出现,不跟在题目后面** —— 否则关掉答案时,
解析会跟着漏出去,而那正是学生不该看到的。
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.schemas.ir import ExamPaper, LessonPlan, Project, QuestionSet

#: 题型 → 卷面上的说法
TYPE_LABELS: dict[str, str] = {
    "mcq": "选择题",
    "short": "简答题",
    "code": "编程题",
    "scenario": "情境题",
}


def _heading(doc: Document, text: str, *, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _paragraph(doc: Document, text: str, *, bold: bool = False, size: int | None = None):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    return paragraph


def _meta_line(doc: Document, project: Project) -> None:
    _paragraph(
        doc,
        f"受众:{project.meta.audience}　课时:{project.meta.duration_min} 分钟",
        size=10,
    )


def _objectives(doc: Document, project: Project) -> None:
    if not project.learning_objectives:
        return
    _heading(doc, "教学目标", level=2)
    for objective in project.learning_objectives:
        doc.add_paragraph(f"{objective.id}　{objective.text}", style="List Bullet")


def _save(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ═══════════════════════════════════════════════════════════════
# 教案
# ═══════════════════════════════════════════════════════════════


def lesson_plan_to_docx(project: Project, plan: LessonPlan) -> bytes:
    """把教案导出成 docx,返回文件字节。"""
    doc = Document()

    title = doc.add_heading(project.meta.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _meta_line(doc, project)
    _objectives(doc, project)

    if plan.overview:
        _heading(doc, "教学设计说明", level=2)
        _paragraph(doc, plan.overview)

    _heading(doc, "教学环节", level=2)

    if plan.timeline:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        header = table.rows[0].cells
        for index, label in enumerate(
            ["环节", "时长", "教师做什么", "学生做什么", "怎么确认听懂了"]
        ):
            header[index].text = label

        for entry in plan.timeline:
            cells = table.add_row().cells
            cells[0].text = entry.phase
            cells[1].text = f"{entry.minutes} 分钟"
            cells[2].text = entry.teacher_action
            cells[3].text = entry.student_action
            cells[4].text = entry.check_understanding

        total = sum(entry.minutes for entry in plan.timeline)
        _paragraph(doc, f"合计:{total} 分钟", bold=True)
    else:
        _paragraph(doc, "(还没有写教学环节)")

    if plan.key_points:
        _heading(doc, "重点", level=2)
        for point in plan.key_points:
            doc.add_paragraph(point, style="List Bullet")

    if plan.difficult_points:
        _heading(doc, "难点", level=2)
        for point in plan.difficult_points:
            doc.add_paragraph(point, style="List Bullet")

    return _save(doc)


# ═══════════════════════════════════════════════════════════════
# 考卷
# ═══════════════════════════════════════════════════════════════


def exam_paper_to_docx(
    project: Project,
    question_set: QuestionSet,
    paper: ExamPaper,
    *,
    include_answer_key: bool = True,
) -> bytes:
    """把模拟卷导出成 docx。

    `include_answer_key=False` 得到学生版:题目在,答案和解析不在。
    """
    doc = Document()

    doc.add_heading(f"{project.meta.title} · 试卷", level=0)
    constraints = paper.constraints
    _paragraph(
        doc,
        f"总分:{constraints.total_points} 分　时长:{constraints.duration_min} 分钟",
        size=10,
    )
    _meta_line(doc, project)

    by_id = {q.id: q for q in question_set.questions}
    ordered: list = []

    for section in paper.sections:
        _heading(doc, section.title, level=2)
        if section.rule:
            _paragraph(doc, section.rule, size=10)

        for number, question_id in enumerate(section.question_ids, start=1):
            question = by_id.get(question_id)
            if question is None:
                # 卷子引用了题库里没有的题 —— 这是 §6.4 会报的错。
                # 导出时如实写出来,而不是静默跳过让页数对不上。
                _paragraph(doc, f"{number}. (题目 {question_id} 不在题库中)")
                continue
            ordered.append(question)

            label = TYPE_LABELS.get(question.type, question.type)
            _paragraph(doc, f"{number}. ({label},{question.points} 分) {question.stem}", bold=True)

            for option in question.options:
                doc.add_paragraph(f"{option.id}. {option.text}", style="List Bullet")

    if include_answer_key and ordered:
        _heading(doc, "参考答案与解析", level=1)
        for number, question in enumerate(ordered, start=1):
            if question.answer:
                _paragraph(doc, f"{number}. 参考答案:{question.answer}", bold=True)
            # 解析只出现在这里 —— 跟在题目后面的话,关掉答案时会一起漏出去
            if question.explanation:
                _paragraph(doc, f"　　解析:{question.explanation}")

    return _save(doc)
