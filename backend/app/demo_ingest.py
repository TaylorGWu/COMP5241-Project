"""材料摄取演示 —— 阶段 1。

    cd backend && ../.venv/bin/python -m app.demo_ingest

**不需要 API key**(OCR 用假实现)。

它会现场造出 7 种材料,把最容易踩的坑都放进去:

    ✅ 有文字层的 PDF        → 直接抽到
    ⚠️  扫描版 PDF(纯图片)  → 抽不到字 → 自动转 OCR
    ✅ PPTX / DOCX / Markdown → 直接抽到
    ⚠️  图片                  → 走 OCR
    ❌ 不支持格式             → 明确拒绝,说明支持哪些

重点看最后那张表:**每份材料都报告抽到多少字**,而不是笼统的「成功」。
"""

from __future__ import annotations

import asyncio
import tempfile
from pathlib import Path

try:
    import pymupdf
except ImportError:  # pragma: no cover
    import fitz as pymupdf  # type: ignore[no-redef]

from app.services.ingest import IngestDeps, IngestReport, ingest_many, total_chunks

#: PyMuPDF 内置的简体中文字体
CJK = "china-s"


# ═══════════════════════════════════════════════════════════════
# 造测试材料
# ═══════════════════════════════════════════════════════════════


def make_text_pdf(path: Path) -> None:
    """有文字层的 PDF —— 就像老师从 Word 导出的那种。"""
    doc = pymupdf.open()
    for i, body in enumerate(
        [
            "Git 分支与合并\n\n分支让不同人的改动互相隔离,主干始终保持可用。\n"
            "合并之前先确认改动可审查,并让至少一位同伴看过。",
            "冲突是怎么产生的\n\n两个人改了同一文件的同一行时会产生冲突。\n"
            "解决办法是先拉取最新代码,再手动决定保留哪一份。",
        ],
        start=1,
    ):
        page = doc.new_page()
        try:
            page.insert_text((60, 90), body, fontname=CJK, fontsize=13)
        except Exception:
            page.insert_text((60, 90), f"Page {i}: Git branching and merging.", fontsize=13)
    doc.save(str(path))
    doc.close()


def make_scanned_pdf(src: Path, dst: Path) -> None:
    """扫描版 PDF —— 把上面那页渲染成图,再塞进新 PDF。

    结果是一份**没有文字层**的 PDF:扩展名一样,但 get_text() 返回空。
    这是最容易被静默忽略的一种材料。
    """
    with pymupdf.open(str(src)) as doc:
        page = doc[0]
        pix = page.get_pixmap(dpi=150)
        out = pymupdf.open()
        new_page = out.new_page(width=page.rect.width, height=page.rect.height)
        new_page.insert_image(new_page.rect, pixmap=pix)
        out.save(str(dst))
        out.close()


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "需求获取:先问流程"
    body = slide.placeholders[1].text_frame
    body.text = "不要一上来就问「要不要做个系统」"
    p = body.add_paragraph()
    p.text = "先问:现在这件事是怎么做的?哪里会卡住?"
    p.font.size = Pt(18)
    slide.notes_slide.notes_text_frame.text = "备注也算材料内容 —— 讲解常写在这里。"
    prs.save(str(path))


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("AI 编码代理的使用要点", 0)
    d.add_paragraph("把任务写成一份简报:目标、约束、验收标准。")
    d.add_paragraph("生成之后要审查输出,而不是直接采用。")
    d.add_paragraph("遇到不确定的地方,让代理说明它的依据。")
    d.save(str(path))


def make_markdown(path: Path) -> None:
    path.write_text(
        "# 版本控制基础\n\n"
        "## 提交信息怎么写\n\n"
        "- 说清「为什么改」,而不只是「改了什么」\n"
        "- 一个提交只做一件事\n"
        "- 关联对应的 issue 编号\n",
        encoding="utf-8",
    )


def make_png(path: Path) -> None:
    """造一张带文字的截图,模拟「老师拍的书页」。"""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 300), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 10, 890, 290], outline="black", width=2)
    draw.text((30, 40), "Textbook page (photo / screenshot)", fill="black")
    draw.text((30, 90), "Branches isolate your changes.", fill="black")
    draw.text((30, 130), "Merge only after review.", fill="black")
    img.save(str(path))


def make_unsupported(path: Path) -> None:
    """一个不支持的二进制文件(伪造的 xlsx)。"""
    path.write_bytes(b"PK\x03\x04" + b"\x00" * 200 + b"xl/worksheets/sheet1.xml")


# ═══════════════════════════════════════════════════════════════
# 假 OCR:让 demo 不需要 key
# ═══════════════════════════════════════════════════════════════


async def fake_ocr(image_bytes: bytes) -> str:
    """假 OCR。特意回显图片大小 —— 这样能确认真的走到了 OCR 这条路。"""
    return (
        "【OCR 识别结果】分支让不同人的改动互相隔离,主干始终保持可用。"
        "合并之前先确认改动可审查。"
        f"(输入图片 {len(image_bytes) // 1024} KB)"
    )


async def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        files: list[Path] = []

        print("═" * 74)
        print("造测试材料")
        print("═" * 74)

        text_pdf = tmpdir / "讲义.pdf"
        make_text_pdf(text_pdf)
        files.append(text_pdf)

        scanned = tmpdir / "扫描版讲义.pdf"
        make_scanned_pdf(text_pdf, scanned)
        files.append(scanned)

        pptx_file = tmpdir / "课件.pptx"
        make_pptx(pptx_file)
        files.append(pptx_file)

        docx_file = tmpdir / "教案.docx"
        make_docx(docx_file)
        files.append(docx_file)

        md_file = tmpdir / "笔记.md"
        make_markdown(md_file)
        files.append(md_file)

        png_file = tmpdir / "书页截图.png"
        make_png(png_file)
        files.append(png_file)

        bad_file = tmpdir / "报表.xlsx"
        make_unsupported(bad_file)
        files.append(bad_file)

        for f in files:
            print(f"  {f.name:<24} {f.stat().st_size:>8,} 字节")

        # ── 摄取 ────────────────────────────────────────────────
        print()
        print("═" * 74)
        print("摄取结果")
        print("═" * 74)

        deps = IngestDeps(ocr_image=fake_ocr)
        reports = await ingest_many(files, deps=deps)

        for r in reports:
            print("  " + r.summary())

        # ── 详情:证明页码真的保留了 ──────────────────────────────
        print()
        print("═" * 74)
        print("分块详情(页码是溯源的基础)")
        print("═" * 74)
        for r in reports:
            if not r.chunks:
                continue
            print(f"\n  【{r.filename}】{len(r.chunks)} 块")
            for c in r.chunks[:3]:
                preview = c.text.replace("\n", " ")[:58]
                print(f"    p.{c.page} [{c.kind}] {preview}…")

        # ── 汇总 ────────────────────────────────────────────────
        print()
        print("═" * 74)
        print("汇总")
        print("═" * 74)
        ok = sum(1 for r in reports if r.status == "ok")
        degraded = sum(1 for r in reports if r.status == "degraded")
        failed = sum(1 for r in reports if r.status == "failed")
        print(f"  ✅ 正常 {ok} 个 · ⚠️  降级 {degraded} 个 · ❌ 失败 {failed} 个")
        print(f"  共抽出 {sum(r.extracted_chars for r in reports):,} 字,"
              f"{total_chunks(reports)} 个检索块")
        print()
        print("  注意「扫描版讲义.pdf」:扩展名和「讲义.pdf」一样,")
        print("  但它的 detected_form 是 scanned_pdf,而且走了 OCR ——")
        print("  这就是「探测」而不是「猜扩展名」的意义。")


if __name__ == "__main__":
    asyncio.run(main())
